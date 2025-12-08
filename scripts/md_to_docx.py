"""
Convert main.md to DOCX format with embedded images.
"""
import re
import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def parse_markdown_to_docx(md_path: str, docx_path: str, base_dir: str):
    """Convert markdown file to DOCX with images."""
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    
    lines = content.split('\n')
    i = 0
    in_table = False
    table_rows = []
    
    while i < len(lines):
        line = lines[i]
        
        # Handle images
        img_match = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', line.strip())
        if img_match:
            caption = img_match.group(1)
            img_path = img_match.group(2)
            full_path = os.path.join(base_dir, img_path)
            if os.path.exists(full_path):
                try:
                    doc.add_picture(full_path, width=Inches(5.5))
                    p = doc.add_paragraph(caption)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.runs[0]
                    run.italic = True
                    run.font.size = Pt(10)
                except Exception as e:
                    doc.add_paragraph(f"[Image: {caption}]")
            else:
                doc.add_paragraph(f"[Image not found: {img_path}]")
            i += 1
            continue
        
        # Handle tables
        if '|' in line and line.strip().startswith('|'):
            if not in_table:
                in_table = True
                table_rows = []
            cells = [c.strip() for c in line.split('|')[1:-1]]
            if not all(c.replace('-', '').replace(':', '') == '' for c in cells):
                table_rows.append(cells)
            i += 1
            continue
        elif in_table:
            if table_rows:
                num_cols = max(len(row) for row in table_rows)
                table = doc.add_table(rows=len(table_rows), cols=num_cols)
                table.style = 'Table Grid'
                for row_idx, row_data in enumerate(table_rows):
                    for col_idx, cell_data in enumerate(row_data):
                        if col_idx < num_cols:
                            cell = table.rows[row_idx].cells[col_idx]
                            text = cell_data.replace('**', '').replace('*', '')
                            cell.text = text
                doc.add_paragraph()
            in_table = False
            table_rows = []
        
        # Skip horizontal rules
        if line.strip() == '---':
            doc.add_paragraph()
            i += 1
            continue
        
        # Handle headings
        if line.startswith('# '):
            p = doc.add_heading(line[2:].strip(), level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith('**') and line.endswith('**') and len(line) > 4:
            p = doc.add_paragraph()
            run = p.add_run(line.strip('*'))
            run.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('*') and line.endswith('*') and not line.startswith('*   '):
            p = doc.add_paragraph()
            run = p.add_run(line.strip('*'))
            run.italic = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('*   ') or line.startswith('-   '):
            text = line[4:].strip()
            text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
            text = re.sub(r'\*([^*]+)\*', r'\1', text)
            text = re.sub(r'`([^`]+)`', r'\1', text)
            p = doc.add_paragraph(text, style='List Bullet')
        elif line.strip().startswith('1.') or line.strip().startswith('2.') or line.strip().startswith('3.') or line.strip().startswith('4.'):
            text = re.sub(r'^\d+\.\s*', '', line.strip())
            text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
            text = re.sub(r'\*([^*]+)\*', r'\1', text)
            text = re.sub(r'`([^`]+)`', r'\1', text)
            doc.add_paragraph(text, style='List Number')
        elif line.strip():
            text = line.strip()
            text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
            text = re.sub(r'\*([^*]+)\*', r'\1', text)
            text = re.sub(r'`([^`]+)`', r'\1', text)
            text = re.sub(r'\$[^$]+\$', '', text)  # Remove math notation
            doc.add_paragraph(text)
        
        i += 1
    
    doc.save(docx_path)
    print(f"Created: {docx_path}")

if __name__ == '__main__':
    base_dir = 'resources/paper'
    parse_markdown_to_docx(
        'resources/paper/main.md',
        'resources/paper/TriArchitect_Submission_Final.docx',
        base_dir
    )
